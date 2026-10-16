import re
import subprocess
from pathlib import Path

import markdown
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
TITLE = "AI-детектив"
LAB_NO = 6
AUTHORS = ["Мещеряков Даниил Павлович, ИСУ 409130", "Соболь Владимир Вячеславович, ИСУ 409594"]
FONT = "Times New Roman"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


def set_font(run, size=None, bold=None, code=False):
    run.font.name = "Consolas" if code else FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), run.font.name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(par, text, size=None):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            set_font(par.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("`"):
            set_font(par.add_run(part[1:-1]), (size - 1) if size else 11, code=True)
        else:
            set_font(par.add_run(part), size)


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, bold=None, before=None, after=None):
    p = doc.add_paragraph()
    p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if bold:
        set_font(p.add_run(text), size, bold=True)
    else:
        add_inline(p, text, size)
    return p


def title_page(doc):
    c = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Федеральное государственное автономное образовательное учреждение высшего образования", c, after=0)
    para(doc, "«Национальный исследовательский университет ИТМО»", c, bold=True, after=0)
    para(doc, "Факультет инфокоммуникационных технологий", c, before=6, after=0)
    para(doc, "ОТЧЕТ", c, size=16, bold=True, before=150, after=0)
    para(doc, f"по лабораторной работе №{LAB_NO}", c, after=0)
    para(doc, f"«{TITLE}»", c, bold=True, after=0)
    para(doc, "по дисциплине «Автоматическая обработка текстов»", c, after=0)
    para(doc, "Выполнили студенты группы К3440:", WD_ALIGN_PARAGRAPH.RIGHT, before=150, after=0)
    for a in AUTHORS:
        para(doc, a, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(doc, "Санкт-Петербург", c, before=150, after=0)
    para(doc, "2026", c, after=0)
    doc.add_page_break()


def shade(cell, color):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_table(doc, lines):
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    plain = [[re.sub(r"[*`]", "", c) for c in r] for r in rows]
    cols = range(len(rows[0]))
    min_w = [max(len(w) for r in plain for w in (r[j].split() or [""])) * 0.2 + 0.4 for j in cols]
    pref = [sum(len(r[j]) for r in plain) / len(plain) + 1 for j in cols]
    extra = max(0.0, 17 - sum(min_w))
    widths = [Cm((m + extra * p / sum(pref)) * min(1, 17 / sum(min_w))) for m, p in zip(min_w, pref)]
    for i, r in enumerate(rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j, text in enumerate(r):
            cell = table.cell(i, j)
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i == 0:
                p.paragraph_format.keep_with_next = True
                set_font(p.add_run(text), 10, bold=True)
                shade(cell, "E7ECF3")
            else:
                add_inline(p, text, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_numbers(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)
    run.font.size = Pt(10)


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 15), ("Heading 2", 13)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), False
        st.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        fonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), FONT)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)


def body_lines(md_path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    start = next(i for i, ln in enumerate(lines) if ln.startswith("## "))
    return lines[start:]


def build_docx(md_path, out_path):
    doc = Document()
    setup(doc)
    title_page(doc)
    page_numbers(doc)
    lines = body_lines(md_path)
    i = 0
    while i < len(lines):
        ln = lines[i]
        img = re.match(r"^!\[(.*)\]\((.+)\)$", ln.strip())
        if ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=1)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:].strip(), level=2)
        elif img:
            doc.add_picture(str(ROOT / img.group(2)), width=Cm(16))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if img.group(1):
                para(doc, img.group(1), WD_ALIGN_PARAGRAPH.CENTER, size=10)
        elif ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif re.match(r"^\s*[-*] ", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\d+\. ", ln):
            p = para(doc, ln.strip())
            p.paragraph_format.left_indent = Cm(0.6)
        elif ln.strip():
            para(doc, ln.strip())
        i += 1
    doc.save(out_path)


CSS = """
@page { size: A4; margin: 2cm; }
body { font-family: 'Times New Roman', serif; font-size: 12pt; line-height: 1.35; color: #111; }
h2 { color: #2E74B5; font-weight: normal; font-size: 15pt; margin: 18pt 0 6pt; }
h3 { color: #2E74B5; font-weight: normal; font-size: 13pt; margin: 12pt 0 4pt; }
p { text-align: justify; margin: 0 0 6pt; }
table { border-collapse: collapse; margin: 6pt auto 10pt; font-size: 10pt; page-break-inside: avoid; }
th, td { border: 1px solid #777; padding: 2pt 5pt; vertical-align: top; }
th { background: #E7ECF3; }
code { font-family: Consolas, monospace; font-size: 10.5pt; }
img { display: block; margin: 6pt auto; max-width: 16cm; }
.title { height: 25.5cm; display: flex; flex-direction: column; text-align: center; page-break-after: always; }
.title .r { text-align: right; }
.grow { flex: 1; }
"""


def spaced_lists(lines):
    out = []
    item = re.compile(r"^(\s*[-*] |\d+\. )")
    for ln in lines:
        if item.match(ln) and out and out[-1].strip() and not item.match(out[-1]):
            out.append("")
        out.append(ln)
    return out


def build_pdf(md_path, out_path):
    body = markdown.markdown("\n".join(spaced_lists(body_lines(md_path))), extensions=["tables"])
    body = body.replace('src="results/', f'src="{(ROOT / "results").as_uri()}/')
    body = re.sub(r'(<img alt="([^"]+)"[^>]*>)', r"\1<p style='text-align:center;font-size:10pt'>\2</p>", body)
    authors = "".join(f"<div class='r'>{a}</div>" for a in AUTHORS)
    title = (f"<div class='title'><div>Федеральное государственное автономное образовательное учреждение высшего "
             f"образования</div><div><b>«Национальный исследовательский университет ИТМО»</b></div>"
             f"<div style='margin-top:6pt'>Факультет инфокоммуникационных технологий</div><div class='grow'></div>"
             f"<div style='font-size:16pt'><b>ОТЧЕТ</b></div><div>по лабораторной работе №{LAB_NO}</div>"
             f"<div><b>«{TITLE}»</b></div><div>по дисциплине «Автоматическая обработка текстов»</div>"
             f"<div class='grow'></div><div class='r'>Выполнили студенты группы К3440:</div>{authors}"
             f"<div class='grow'></div><div>Санкт-Петербург</div><div>2026</div></div>")
    html = f"<!doctype html><html lang='ru'><head><meta charset='utf-8'><style>{CSS}</style></head><body>{title}{body}</body></html>"
    html_path = ROOT / "results" / "report.html"
    html_path.write_text(html, encoding="utf-8")
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={out_path}", html_path.as_uri()], check=True, capture_output=True)


def count_words(md_path):
    text = "\n".join(ln for ln in body_lines(md_path) if not ln.startswith("!["))
    text = re.sub(r"[|*`#]", " ", text)
    return len([w for w in text.split() if re.search(r"\w", w)])


if __name__ == "__main__":
    md = ROOT / "report.md"
    build_docx(md, ROOT / "report.docx")
    build_pdf(md, ROOT / "report.pdf")
    print(f"report.docx и report.pdf собраны, слов в отчете: {count_words(md)}")
