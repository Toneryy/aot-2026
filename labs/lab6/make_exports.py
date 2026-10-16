import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.dossier import DATA, RESULTS, load_claims, load_documents, load_questions, read_jsonl
from src.stress import load_contradictions, load_removals

ROOT = Path(__file__).parent
PREFIX = "ЛР6_К3440_Мещеряков_Соболь"


def style(ws, widths=None):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E7ECF3")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(ws.columns, 1):
        values = [str(c.value) if c.value is not None else "" for c in col]
        width = min(60, max(8, max(len(v) for v in values[:200]) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
        for c in col[1:]:
            c.alignment = Alignment(wrap_text=width >= 60, vertical="top")
    ws.freeze_panes = "A2"


def to_xlsx(path, sheets):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            style(xw.sheets[name[:31]])


def dossier_frames():
    docs = pd.DataFrame([{**d.model_dump(), "words": len(d.text.split())} for d in load_documents()])
    claims = pd.DataFrame([{**c.model_dump(), "padic": str(list(c.padic)), "derived_from": " ".join(c.derived_from)}
                           for c in load_claims()])
    qs = pd.DataFrame([{**q.model_dump(), "gold_evidence_ids": " ".join(q.gold_evidence_ids),
                        "gold_counterevidence_ids": " ".join(q.gold_counterevidence_ids)} for q in load_questions()])
    rem = pd.DataFrame(load_removals())
    con = pd.DataFrame([{"variant_id": s["variant_id"], "question_id": s["question_id"],
                         "expected_status": s["expected_status"], "document_id": s["document"]["document_id"],
                         "document_title": s["document"]["title"], "claim_id": s["claim"]["claim_id"],
                         "claim_text": s["claim"]["text"], "evidence": s["claim"]["evidence"], "check": s["check"]}
                        for s in load_contradictions()])
    return {"Документы": docs, "Утверждения": claims, "Вопросы и эталоны": qs, "Удаление посылки": rem,
            "Противоречия": con}


def results_frames():
    names = {"comparison": "Три режима по вопросам", "metrics": "Метрики по режимам", "stress": "Стресс-тесты",
             "stability": "Устойчивость", "calibration": "Калибровка", "flat_vs_padic": "Плоский и p-adic",
             "retrieval": "Поиск BM25"}
    return {title: pd.read_csv(RESULTS / f"{name}.csv") for name, title in names.items()}


def main(out_dir):
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    to_xlsx(out / f"{PREFIX}_досье_и_эталоны.xlsx", dossier_frames())
    to_xlsx(out / f"{PREFIX}_результаты.xlsx", results_frames())
    comp = pd.read_csv(RESULTS / "comparison.csv")
    comp.to_csv(out / f"{PREFIX}_три_режима.csv", index=False, encoding="utf-8-sig")
    to_xlsx(out / f"{PREFIX}_три_режима.xlsx", {"Три режима": comp})
    st = pd.read_csv(RESULTS / "stress.csv")
    st.to_csv(out / f"{PREFIX}_стресс_тесты.csv", index=False, encoding="utf-8-sig")
    to_xlsx(out / f"{PREFIX}_стресс_тесты.xlsx", {"Стресс-тесты": st})
    mm = pd.read_csv(RESULTS / "metrics.csv")
    mm.to_csv(out / f"{PREFIX}_метрики.csv", index=False, encoding="utf-8-sig")
    to_xlsx(out / f"{PREFIX}_метрики.xlsx", {"Метрики": mm, "Устойчивость": pd.read_csv(RESULTS / "stability.csv")})
    shutil.copy(ROOT / "report.docx", out / f"{PREFIX}_отчет.docx")
    shutil.copy(ROOT / "report.pdf", out / f"{PREFIX}_отчет.pdf")
    shutil.copy(ROOT / "AI_LOG.md", out / f"{PREFIX}_AI_LOG.md")
    shutil.copy(ROOT / "CONTRIBUTIONS.md", out / f"{PREFIX}_вклад.md")
    shutil.copy(ROOT / "prompts" / "system.txt", out / f"{PREFIX}_промпт.txt")
    for name in ("documents", "claims", "questions"):
        shutil.copy(DATA / f"{name}.jsonl", out / f"{PREFIX}_{name}.jsonl")
    print("\n".join(sorted(p.name for p in out.iterdir())))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ROOT / "на отправку ЛР6")
