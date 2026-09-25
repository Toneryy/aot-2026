import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
TITLE = "Мини-RAG с доказательствами"
LAB_NO = 5
AUTHORS = ["Мещеряков Даниил Павлович, ИСУ 409130", "Соболь Владимир Вячеславович, ИСУ 409594"]
TEACHER = "Гусарова Наталия Федоровна"
FONT = "Times New Roman"


def set_font(run, size=None, bold=None, code=False):
    run.font.name = "Consolas" if code else FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), run.font.name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(par, text, size=None):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            set_font(par.add_run(part[2:-2]), size, bold=True)
        elif part.startswith("`"):
            set_font(par.add_run(part[1:-1]), size - 1 if size else 11, code=True)
        else:
            set_font(par.add_run(part), size)


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, bold=None, before=None, after=None):
    p = doc.add_paragraph()
    p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if bold:
        set_font(p.add_run(text), size, bold=True)
    else:
        add_inline(p, text, size)
    return p


def title_page(doc):
    c = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, "Федеральное государственное автономное образовательное учреждение высшего образования", c, after=0)
    para(doc, "«Национальный исследовательский университет ИТМО»", c, bold=True, after=0)
    para(doc, "Факультет инфокоммуникационных технологий", c, before=6, after=0)
    para(doc, "ОТЧЕТ", c, size=16, bold=True, before=150, after=0)
    para(doc, f"по лабораторной работе №{LAB_NO}", c, after=0)
    para(doc, f"«{TITLE}»", c, bold=True, after=0)
    para(doc, "по дисциплине «Автоматическая обработка текстов»", c, after=0)
    para(doc, "Выполнили студенты группы К3440:", WD_ALIGN_PARAGRAPH.RIGHT, before=120, after=0)
    for a in AUTHORS:
        para(doc, a, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(doc, "Преподаватель:", WD_ALIGN_PARAGRAPH.RIGHT, before=12, after=0)
    para(doc, TEACHER, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(doc, "Санкт-Петербург", c, before=110, after=0)
    para(doc, "2026", c, after=0)
    doc.add_page_break()


def shade(cell, color):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_table(doc, lines):
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    plain = [[re.sub(r"[*`]", "", c) for c in r] for r in rows]
    cols = range(len(rows[0]))
    min_w = [max(len(w) for r in plain for w in (r[j].split() or [""])) * 0.21 + 0.4 for j in cols]
    pref = [sum(len(r[j]) for r in plain) / len(plain) + 1 for j in cols]
    extra = max(0.0, 17 - sum(min_w))
    widths = [Cm((m + extra * p / sum(pref)) * min(1, 17 / sum(min_w))) for m, p in zip(min_w, pref)]
    for i, r in enumerate(rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for j, text in enumerate(r):
            cell = table.cell(i, j)
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if i == 0:
                p.paragraph_format.keep_with_next = True
                set_font(p.add_run(text), 10.5, bold=True)
                shade(cell, "E7ECF3")
            else:
                add_inline(p, text, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_numbers(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)
    run.font.size = Pt(10)


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = FONT, Pt(size), False
        st.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        fonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), FONT)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)


def build(md_path, out_path):
    doc = Document()
    setup(doc)
    title_page(doc)
    page_numbers(doc)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    start = next(i for i, ln in enumerate(lines) if ln.startswith("## "))
    i = start
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=1)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:].strip(), level=2)
        elif ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        elif re.match(r"^\s*[-*] ", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\d+\. ", ln):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", ln))
        elif ln.strip():
            para(doc, ln.strip())
        i += 1
    doc.save(out_path)


if __name__ == "__main__":
    build(ROOT / "report.md", ROOT / "report.docx")
    print("report.docx written")
